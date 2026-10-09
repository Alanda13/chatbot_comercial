"""
O Excel do botão "⬇️ Baixar em Excel", no formato institucional (o mesmo
do resumo gerencial do supervisor): faixa azul com a logo, título e
subtítulo; cabeçalho azul; linhas zebradas; largura automática; cabeçalho
fixo ao rolar.

A tabela chega já formatada pra tela ("R$ 1.234,56", "2,29%"): aqui esses
textos voltam a ser NÚMEROS, com o formato do Excel, pra dar pra somar.
"""
import io
import math
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagemExcel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

from config.constants import AZUL_FERRONORTE

_AZUL = AZUL_FERRONORTE.lstrip("#")
_ZEBRA = "E8F1F9"

_FUNDO_AZUL = PatternFill("solid", fgColor=_AZUL)
_FUNDO_BRANCO = PatternFill("solid", fgColor="FFFFFF")
_FUNDO_ZEBRA = PatternFill("solid", fgColor=_ZEBRA)
_LINHA = Side(style="thin", color="BFBFBF")
_BORDA = Border(left=_LINHA, right=_LINHA, top=_LINHA, bottom=_LINHA)

_LINHA_CABECALHO = 5  # 1: logo; 2-3: faixa com pergunta e subtítulo; 4: em branco
_ALTURA_LOGO = 60  # px
_LARGURA_MINIMA = 45  # caracteres — largura mínima da tabela (e da faixa)

# Textos da tela → número + formato do Excel. O sinal explícito ("+R$",
# "+5,64%") é das colunas de variação e continua aparecendo no Excel.
_MOEDA = re.compile(r"^([+−-]?)R\$ ([\d.]+,\d+)$")
_PERCENTUAL = re.compile(r"^([+−-]?)([\d.]+,\d+)%$")
_NUMERO = re.compile(r"^([+−-]?)(\d{1,3}(?:\.\d{3})*(?:,\d+)?)$")
_SEM_VALOR = {"sem dados", "menos de 0,01%"}


def _numero(texto):
    return float(texto.replace(".", "").replace(",", "."))


def _formato(base, sinal):
    if sinal == "+":
        return f"+{base};-{base};{base}"
    return base


def _converter(texto):
    """'R$ 1.234,56' → (1234.56, formato); texto que não é valor → None."""
    escala = 1

    if achado := _MOEDA.match(texto):
        sinal, valor = achado.groups()
        formato = '"R$" #,##0.00'
    elif achado := _PERCENTUAL.match(texto):
        sinal, valor = achado.groups()
        formato, escala = "0.00%", 100
    elif achado := _NUMERO.match(texto):
        sinal, valor = achado.groups()
        formato = "#,##0.00" if "," in valor else "#,##0"
    else:
        return None

    numero = _numero(valor) / escala
    if sinal in ("−", "-"):
        numero = -numero

    return numero, _formato(formato, sinal)


def _colunas_de_valor(tabela):
    """Colunas de texto em que toda célula é valor (R$, %, número) ou
    "sem dados" — as de nome/filial/mês ficam como estão."""
    colunas = []

    for coluna in tabela.columns:
        celulas = [c for c in tabela[coluna] if c not in _SEM_VALOR]

        if celulas and all(isinstance(c, str) and _converter(c) for c in celulas):
            colunas.append(coluna)

    return colunas


def _colocar_logo(planilha, caminho_logo):
    largura, altura = Image.open(caminho_logo).size
    imagem = ImagemExcel(caminho_logo)
    imagem.height = _ALTURA_LOGO
    imagem.width = round(largura * _ALTURA_LOGO / altura)
    planilha.add_image(imagem, "A1")


def _largura(planilha, coluna):
    letra = get_column_letter(coluna)
    # Coluna sem largura definida: a padrão do Excel (8,43 caracteres).
    if letra not in planilha.column_dimensions:
        return 8.43
    return planilha.column_dimensions[letra].width


def _montar_faixa(planilha, titulo, subtitulo, caminho_logo, colunas_da_tabela):
    """Topo do tamanho da tabela: a logo na 1ª linha (em fundo branco — as
    letras azuis sumiam no azul) e, embaixo, a faixa azul com a pergunta e
    o subtítulo. A pergunta que não cabe na largura da tabela quebra em
    mais linhas (a faixa passando da tabela ficava estranha)."""
    if caminho_logo:
        _colocar_logo(planilha, caminho_logo)
    planilha.row_dimensions[1].height = _ALTURA_LOGO * 0.75 + 6  # px → pt

    # Tabela estreita (1 ou 2 colunas): alarga a última pra pergunta não
    # virar uma coluna de palavras.
    largura_tabela = sum(_largura(planilha, c) for c in range(1, colunas_da_tabela + 1))
    if largura_tabela < _LARGURA_MINIMA:
        letra = get_column_letter(colunas_da_tabela)
        planilha.column_dimensions[letra].width = (
            _largura(planilha, colunas_da_tabela) + _LARGURA_MINIMA - largura_tabela
        )
        largura_tabela = _LARGURA_MINIMA

    textos = [
        (2, titulo, Font(bold=True, size=14, color="FFFFFF"), 1.2, 19),
        (3, subtitulo, Font(italic=True, size=10, color=_ZEBRA), 1.0, 14),
    ]

    for linha, texto, fonte, largura_letra, altura_linha in textos:
        planilha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=colunas_da_tabela)
        for coluna in range(1, colunas_da_tabela + 1):
            planilha.cell(linha, coluna).fill = _FUNDO_AZUL

        celula = planilha.cell(linha, 1, texto or "")
        celula.font = fonte
        celula.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

        # O Excel não ajusta a altura de célula mesclada sozinho.
        linhas = math.ceil(len(texto or "") * largura_letra / (largura_tabela - 2)) or 1
        planilha.row_dimensions[linha].height = linhas * altura_linha + 8


def gerar_excel(tabela: pd.DataFrame, titulo, subtitulo=None, caminho_logo=None) -> bytes:
    livro = Workbook()
    planilha = livro.active
    planilha.title = "Dados"

    colunas = list(tabela.columns)

    for indice, nome in enumerate(colunas, start=1):
        celula = planilha.cell(_LINHA_CABECALHO, indice, str(nome))
        celula.font = Font(bold=True, size=11, color="FFFFFF")
        celula.fill = _FUNDO_AZUL
        celula.border = _BORDA
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    de_valor = set(_colunas_de_valor(tabela))
    larguras = [len(str(nome)) for nome in colunas]

    for posicao, valores in enumerate(tabela.itertuples(index=False), start=1):
        linha = _LINHA_CABECALHO + posicao

        for indice, (nome, valor) in enumerate(zip(colunas, valores), start=1):
            texto = "" if valor is None or (isinstance(valor, float) and pd.isna(valor)) else str(valor)
            larguras[indice - 1] = max(larguras[indice - 1], len(texto))

            celula = planilha.cell(linha, indice)
            celula.border = _BORDA
            celula.font = Font(size=10)

            if posicao % 2 == 0:
                celula.fill = _FUNDO_ZEBRA

            convertido = _converter(texto) if nome in de_valor else None

            if convertido:
                celula.value, celula.number_format = convertido
                celula.alignment = Alignment(horizontal="right")
            else:
                celula.value = valor if isinstance(valor, (int, float)) and texto else texto
                if nome in de_valor:
                    celula.alignment = Alignment(horizontal="right")

    for indice, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = min(largura + 3, 50)

    _montar_faixa(planilha, titulo, subtitulo, caminho_logo, len(colunas))

    planilha.freeze_panes = planilha.cell(_LINHA_CABECALHO + 1, 1)

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
