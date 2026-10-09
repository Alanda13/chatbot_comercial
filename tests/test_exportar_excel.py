import io

import pandas as pd
from openpyxl import load_workbook

from config.settings import CAMINHO_LOGO_SIDEBAR
from ui.exportar_excel import gerar_excel


def _abrir(tabela, **kwargs):
    conteudo = gerar_excel(tabela, titulo="Qual o faturamento?", **kwargs)
    return load_workbook(io.BytesIO(conteudo)).active


def _abrir_com(tabela, titulo):
    return load_workbook(io.BytesIO(gerar_excel(tabela, titulo=titulo))).active


def test_valores_da_tela_viram_numeros_no_excel():
    tabela = pd.DataFrame({
        "Filial": ["TIMON", "JOÃO XXIII"],
        "Faturamento": ["R$ 3.831.746,29", "sem dados"],
        "% Desconto": ["2,29%", "menos de 0,01%"],
        "Variação": ["+R$ 1.500,00", "−R$ 963,53"],
        "Clientes": ["1.234", "15"],
    })

    planilha = _abrir(tabela)

    assert [c.value for c in planilha[5]] == list(tabela.columns)
    assert planilha["B6"].value == 3831746.29
    assert planilha["B6"].number_format == '"R$" #,##0.00'
    assert planilha["B7"].value == "sem dados"
    assert planilha["C6"].value == 0.0229
    assert planilha["C6"].number_format == "0.00%"
    assert planilha["C7"].value == "menos de 0,01%"
    assert planilha["D6"].value == 1500
    assert planilha["D6"].number_format.startswith("+")
    assert planilha["D7"].value == -963.53
    assert planilha["E6"].value == 1234
    # Nome continua texto, mesmo sendo a 1ª coluna.
    assert planilha["A6"].value == "TIMON"


def test_coluna_de_texto_nao_vira_numero():
    tabela = pd.DataFrame({"Mês": ["Janeiro", "2025"], "Faturamento": ["R$ 1,00", "R$ 2,00"]})

    planilha = _abrir(tabela)

    assert planilha["A7"].value == "2025"


def test_titulo_subtitulo_logo_e_cabecalho_fixo():
    tabela = pd.DataFrame({"Filial": ["TIMON"], "Faturamento": ["R$ 1,00"]})

    planilha = _abrir(tabela, subtitulo="Ano: 2026", caminho_logo=CAMINHO_LOGO_SIDEBAR)

    assert planilha["A2"].value == "Qual o faturamento?"
    assert planilha["A3"].value == "Ano: 2026"
    assert len(planilha._images) == 1
    assert planilha.freeze_panes == "A6"
    assert planilha["A5"].font.size == 11


def test_faixa_do_tamanho_da_tabela_e_pergunta_longa_quebra_linha():
    tabela = pd.DataFrame({"RCA": ["LUCIVANDA DA SILVA COSTA"], "Faturamento": ["R$ 1,00"], "% Desconto": ["2,75%"]})
    pergunta = "liste os 5 rcas que mais faturaram mes passado e quanto de desconto cada um deu"

    planilha = _abrir_com(tabela, pergunta)

    assert "A2:C2" in {str(r) for r in planilha.merged_cells.ranges}
    assert planilha.row_dimensions[2].height > 30
